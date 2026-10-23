"""Explicit packaged offline self-test: no database connections or writes."""
from pathlib import Path
import json
import os
import tempfile


def check_authorization(folder):
    """Exercise bundled crypto using a disposable key and synthetic credentials."""
    import base64
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import padding, rsa
    from .authorization import create_authorization
    from .configuration import TPlusApiConfig
    key = rsa.generate_private_key(public_exponent=65537, key_size=1024)
    with tempfile.TemporaryDirectory(prefix="signing-", dir=str(folder)) as directory:
        pem = Path(directory) / "cjet_pri.pem"
        pem.write_bytes(key.private_bytes(serialization.Encoding.PEM,
                                         serialization.PrivateFormat.TraditionalOpenSSL,
                                         serialization.NoEncryption()))
        config = TPlusApiConfig(app_key="offline-app", app_secret="offline-secret", pem_path=str(pem))
        for token in (None, "offline-token"):
            value = create_authorization(config, token, use_sample=False)
            envelope = json.loads(base64.b64decode(value))
            header, payload, signature = envelope["authInfo"].split(".")
            decode = lambda part: base64.urlsafe_b64decode(part + "=" * (-len(part) % 4))
            key.public_key().verify(decode(signature), (header + "." + payload).encode("ascii"),
                                    padding.PSS(mgf=padding.MGF1(hashes.SHA256()), salt_length=32),
                                    hashes.SHA256())
            claims = json.loads(decode(payload))
            assert claims["sub"] == ("chanjet" if token else "tester")
            if token:
                assert claims["access_token"] == token
    return True


def run(folder):
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    import xlrd
    import pyodbc
    from openpyxl import load_workbook
    from PyQt5.QtWidgets import QApplication
    from PyQt5.QtGui import QFont, QFontDatabase
    from .gui import ImportWindow
    from .domain import SourceRow, money_parts
    from .files import export_report
    from decimal import Decimal
    folder = Path(folder).resolve()
    folder.mkdir(parents=True, exist_ok=True)
    application = QApplication.instance() or QApplication([])
    font_id = QFontDatabase.addApplicationFont("C:/Windows/Fonts/msyh.ttc")
    if font_id >= 0:
        application.setFont(QFont(QFontDatabase.applicationFontFamilies(font_id)[0], 9))
    window = ImportWindow()
    rows = [SourceRow(n + 2, "sheet1", {"用友单号": "DEMO" + str(n), "客户名称": "示例医药有限公司", "单据日期": "2026-10-12", "客户编码": "0001", "商品编码": "0099",
            "数量": "10", "含税金额": "67.995"}) for n in range(15)]
    window.model.set_rows(rows, list(rows[0].values))
    window.source_label.setText("离线界面验证：示例数据，未连接数据库")
    window.update_summary()
    window.show()
    application.processEvents()
    screenshot = folder / "ui_preview.png"
    if not window.grab().save(str(screenshot)):
        raise RuntimeError("Offscreen screenshot save failed")
    window.close()
    rows[0].status = "校验失败"
    rows[0].reasons = ["离线测试数据"]
    for row in rows[1:]:
        row.status = "导入成功"
    report = export_report(rows, list(rows[0].values), folder / "demo.xls", "offline", "smoke")
    workbook = load_workbook(report, read_only=True)
    assert "汇总" in workbook.sheetnames
    workbook.close()
    assert money_parts("-67.995", Decimal("0.13")) == (Decimal("-67.995"), Decimal("-60.17"), Decimal("-7.825"))
    result = dict(passed=True, database_connections=0, database_writes=0, xlrd=xlrd.__version__,
                  authorization_signing=check_authorization(folder),
                  report=str(report), screenshot=str(screenshot), odbc_driver_count=len(pyodbc.drivers()))
    (folder / "smoke_result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    return result
