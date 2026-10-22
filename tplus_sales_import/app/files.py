"""BIFF input, atomic reports and durable task snapshots; never import application UI."""
from collections import Counter
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import json
import os
import re
import tempfile
import uuid
import time
import logging
from log_util import logger, diagnostics

from .domain import REQUIRED_COLUMNS, SourceRow


def read_xls(path, stop, progress):
    import xlrd
    started = time.perf_counter()
    source = Path(path)
    logger.info("开始加载XLS file=%s", source)
    if source.suffix.lower() != ".xls":
        raise ValueError("请选择旧式Excel .xls文件")
    book = xlrd.open_workbook(str(source), on_demand=True, formatting_info=True)
    rows, columns = [], []
    try:
        for sheet in book.sheets():
            if stop.is_set():
                raise ValueError("文件加载已停止")
            if not sheet.nrows:
                continue
            headers = [str(sheet.cell_value(0, i)).strip() for i in range(sheet.ncols)]
            if not any(headers):
                continue
            missing = set(REQUIRED_COLUMNS) - set(headers)
            if missing:
                raise ValueError(f"工作表{sheet.name}缺列：" + "、".join(sorted(missing)))
            nonempty = [name for name in headers if name]
            if len(nonempty) != len(set(nonempty)):
                raise ValueError(f"工作表{sheet.name}存在重复列名")
            columns.extend(name for name in nonempty if name not in columns)
            for index in range(1, sheet.nrows):
                if stop.is_set():
                    raise ValueError("文件加载已停止")
                values = {}
                for col, name in enumerate(headers):
                    if not name:
                        continue
                    cell = sheet.cell(index, col)
                    value = cell.value
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        raise ValueError(f"{sheet.name}第{index + 1}行{name}为Excel错误值")
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        value = format(Decimal(str(value)), "f")
                        if Decimal(value) == Decimal(value).to_integral():
                            value = str(int(Decimal(value)))
                            if name.endswith("编码") or name in ("用友单号", "销售标识"):
                                xf = book.xf_list[cell.xf_index]
                                fmt = book.format_map[xf.format_key].format_str
                                if re.fullmatch(r"0+", fmt):
                                    value = value.zfill(len(fmt))
                    values[name] = value
                if any(str(value).strip() for value in values.values()):
                    rows.append(SourceRow(index + 1, sheet.name, values))
                if index % 500 == 0:
                    progress("加载 " + sheet.name, index, sheet.nrows - 1)
        if not rows:
            raise ValueError("文件没有可导入的明细")
        logger.info("XLS加载完成 file=%s sheets=%d rows=%d documents=%d elapsed_seconds=%.3f", source, book.nsheets, len(rows), len({row.voucher for row in rows if row.voucher}), time.perf_counter() - started)
        progress("文件加载完成", len(rows), len(rows))
        return rows, columns
    finally:
        book.release_resources()


def assert_output_writable(source):
    folder = Path(source).resolve().parent
    with tempfile.NamedTemporaryFile(prefix=".tplus_write_check_", dir=str(folder)):
        pass
    return folder


def safe_filename(value):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value))[:80].strip(" .") or "result"


def snapshot_rows(rows):
    return [{"number": row.number, "sheet": row.sheet, "values": row.values,
             "status": row.status, "reasons": row.reasons, "notes": row.notes,
             "sa_code": row.sa_code, "header_id": row.header_id, "parsed": row.parsed} for row in rows]


class TaskJournal:
    def __init__(self, folder, task_id, source, database):
        self.path = Path(folder) / (task_id + ".json")
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.task_id, self.source, self.database = task_id, str(source), database

    def save(self, rows, phase, message=""):
        payload = dict(task_id=self.task_id, source=self.source, database=self.database,
                       phase=phase, message=message, updated=datetime.now().isoformat(), rows=snapshot_rows(rows))
        atomic_json(self.path, payload)

    @staticmethod
    def load(path):
        payload = json.loads(Path(path).read_text(encoding="utf-8"))
        return payload, [SourceRow(**item) for item in payload["rows"]]


def atomic_json(path, payload):
    temporary = Path(str(path) + "." + uuid.uuid4().hex + ".tmp")
    try:
        with temporary.open("w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, default=str)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(str(temporary), str(path))
    finally:
        if temporary.exists():
            temporary.unlink()


def _append_text_safe(sheet, values):
    from openpyxl.cell import WriteOnlyCell
    cells = []
    for value in values:
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False, default=str)
        if value is None:
            value = ""
        # Every string is a literal, even if it starts with =, +, - or @.
        if isinstance(value, str):
            value = "".join(c for c in value if ord(c) >= 32 or c in "\t\n\r")
            cell = WriteOnlyCell(sheet, value=value[:32767])
            cell.data_type = "s"
            cells.append(cell)
        else:
            cells.append(value)
    sheet.append(cells)


def export_report(rows, columns, source, database, task_id, folder=None, message="", include_all=False):
    from openpyxl import Workbook
    started = time.perf_counter()
    task_log = logging.LoggerAdapter(logger, {"task_id": task_id})
    folder = Path(folder) if folder else Path(source).resolve().parent
    folder.mkdir(parents=True, exist_ok=True)
    name = safe_filename(Path(source).stem) + "_" + safe_filename(database) + "_" + safe_filename(task_id) + "_异常结果.xlsx"
    path = folder / name
    if path.exists():
        path = folder / (path.stem + "_" + uuid.uuid4().hex[:8] + path.suffix)
    # A summary is always exported, including when every row succeeds.
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("汇总")
    for record in (("任务", task_id), ("账套", database), ("输入文件", str(source)),
                   ("任务说明", message), ("明细总数", len(rows)),
                   ("源单据数", len({row.voucher for row in rows if row.voucher}))):
        _append_text_safe(summary, record)
    for status, count in sorted(Counter(row.status for row in rows).items()):
        _append_text_safe(summary, ("明细/" + status, count))
    documents = {}
    for index, row in enumerate(rows):
        documents.setdefault(row.voucher or f"空单号/行{index}", set()).add(row.status)
    priority = ("撤销待确认", "待确认", "写入失败", "审批失败", "校验失败", "整单受阻", "未处理", "已撤销", "重复跳过", "导入成功")
    doc_counts = Counter(next((state for state in priority if state in states), "未处理") for states in documents.values())
    for status, count in sorted(doc_counts.items()):
        _append_text_safe(summary, ("单据/" + status, count))
    sheets = {name: workbook.create_sheet(name) for name in ("失败及未处理", "重复跳过", "成功但有说明")}
    if include_all:
        sheets["全部结果"] = workbook.create_sheet("全部结果")
    headers = columns + ["源工作表", "物理行号", "结果状态", "全部原因", "非阻断说明", "SA编号", "表头ID", "客户ID", "存货ID", "单位ID", "仓库ID", "项目ID", "税率比例", "写入含税金额", "写入无税金额", "写入税额"]
    for sheet in sheets.values():
        _append_text_safe(sheet, headers)
    for row in rows:
        parsed = row.parsed
        values = [row.values.get(name, "") for name in columns] + [row.sheet, row.number, row.status,
                  "；".join(row.reasons), "；".join(row.notes), row.sa_code, row.header_id,
                  parsed.get("customers", {}).get("id"), parsed.get("inventory", {}).get("id"),
                  parsed.get("unit", {}).get("id"), parsed.get("warehouse", {}).get("id"),
                  parsed.get("project_id"), str(parsed.get("rate", "")), str(parsed.get("gross", "")),
                  str(parsed.get("net", "")), str(parsed.get("tax", ""))]
        target = "重复跳过" if row.status == "重复跳过" else ("成功但有说明" if row.status == "导入成功" else "失败及未处理")
        if row.status != "导入成功" or row.notes:
            _append_text_safe(sheets[target], values)
        if include_all:
            _append_text_safe(sheets["全部结果"], values)
    temporary = path.with_name("." + path.stem + "." + uuid.uuid4().hex + ".tmp")
    try:
        workbook.save(str(temporary))
        os.replace(str(temporary), str(path))
        task_log.info("结果文件导出成功 path=%s rows=%d marked_rows=%d elapsed_seconds=%.3f", path, len(rows), sum(row.status != "导入成功" or bool(row.notes) for row in rows), time.perf_counter() - started)
    except Exception as exc:
        task_log.error("结果文件导出失败 path=%s %s", path, diagnostics(exc))
        raise
    finally:
        workbook.close()
        if temporary.exists():
            temporary.unlink()
    return path
