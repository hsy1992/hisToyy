from copy import deepcopy
from datetime import date
from decimal import Decimal
from pathlib import Path
import json
import threading

import pytest
from openpyxl import load_workbook

from app.configuration import DatabaseConfig
from app.database import CommitUnknown, DocumentError, PreflightError, SqlRepository
from app.domain import MasterCache, SourceRow, build_payload, decimal_text, key, money_parts, parse_rate, validate_rows
from app.engine import ImportEngine
from app.files import TaskJournal, export_report


def config():
    return DatabaseConfig("fake-host", 1433, "fake-account", "fake-user", "do-not-log-me")


def source_row(source="XKP0026202610122110", number=2, **changes):
    values = {"仓库编码": "线上库", "单据日期": "2026-10-12", "用友单号": source,
              "销售标识": source + str(number), "业务类型": "15", "客户编码": "0001", "项目编码": "送货",
              "商品编码": "0099", "数量": "10", "单价": "6.7995", "含税金额": "67.995"}
    values.update(changes)
    return SourceRow(number, "sheet1", values)


def cache():
    return MasterCache(dict(customers=[dict(id=20, code="0001", name="客户", disabled=0, taxRate=351)],
        inventory=[dict(id=30, code="0099", name="商品", disabled=0, idunit=23, isSingleUnit=1)],
        unit=[dict(id=23, code="23", name="盒", disabled=0)],
        warehouse=[dict(id=15, code="线上库", name="线上库", disabled=0)],
        business=[dict(id=65, code="15", disabled=0), dict(id=66, code="16", disabled=0)],
        project=[dict(id=6, code="送货", name="送货", disabled=0)],
        tax=[dict(id=351, name="13%")]))


class FakeRepository:
    def __init__(self):
        self.committed, self.calls = {}, []
        self.cache = cache()
        self.fail_source = None
        self.stop_after = None
        self.preflight_error = None
        self.unknown = None
        self.closed = False

    def open(self): return "fake-account"
    def close(self): self.closed = True
    def reconnect(self): pass
    def load_cache(self, stop, progress): return self.cache
    def preflight(self, *args):
        if self.preflight_error: raise PreflightError(self.preflight_error)
    def validate_payload(self, *args): pass
    def find_existing(self, sources):
        return {key(s): [self.committed[key(s)]] for s in sources if key(s) in self.committed}
    def max_sequence(self, day): return 77
    def classify(self, exc): return SqlRepository.classify(exc)
    def verify_document(self, doc, identifier): return self.unknown != "corrupt"

    def write_batch(self, documents):
        self.calls.append([doc.source for doc in documents])
        if self.fail_source in [doc.source for doc in documents]:
            raise DocumentError("模拟某单据字段错误，整批已回滚")
        if self.unknown == "absent":
            raise CommitUnknown()
        mapping = {}
        for doc in documents:
            identifier = len(self.committed) + 100
            self.committed[key(doc.source)] = dict(id=identifier, code=doc.code)
            mapping[key(doc.source)] = identifier
        if self.stop_after:
            self.stop_after.set()
        if self.unknown:
            raise CommitUnknown()
        return mapping


def run_engine(tmp_path, rows, repo=None, stop=None, limit=5000):
    source = tmp_path / "source.xls"
    source.touch()
    repository = repo or FakeRepository()
    engine = ImportEngine(config(), source, rows, list(rows[0].values), tmp_path / "runtime",
                          stop or threading.Event(), lambda *args: None,
                          repository_factory=lambda cfg: repository, batch_limit=limit)
    return engine.run(), repository


@pytest.mark.parametrize("raw,expected", [("67.995", ("67.995", "60.17", "7.825")),
    ("-67.995", ("-67.995", "-60.17", "-7.825")), ("0", ("0.00", "0.00", "0.00")),
    ("0.019", ("0.019", "0.02", "-0.001")), ("-0.019", ("-0.019", "-0.02", "0.001"))])
def test_source_gross_preserved_net_half_up_and_symmetric_signs(raw, expected):
    parts = money_parts(raw, Decimal("0.13"))
    assert parts == tuple(map(Decimal, expected))
    assert parts[1] + parts[2] == parts[0]


def test_rate_requires_real_label_not_enum_id():
    assert parse_rate("13%") == parse_rate("0.13") == Decimal("0.13")
    assert parse_rate("0%") == 0
    for value in (None, "", "351", "101", "免税", "NaN", "-1%", "-13", "Infinity"):
        with pytest.raises(ValueError): parse_rate(value)


@pytest.mark.parametrize("label,expected", [("13", "0.13"), ("17", "0.17"),
    ("13.00", "0.13"), (" 13 ", "0.13"), ("0.13", "0.13"),
    ("13%", "0.13"), ("0", "0"), ("1", "0.01"), ("1%", "0.01"),
    ("100", "1"), ("100%", "1")])
def test_customer_enum_name_normalizes_to_ratio_in_payload(label, expected):
    masters = cache()
    masters.enums[351]["name"] = label
    row = source_row()
    doc = validate_rows([row], masters, threading.Event(), lambda *args: None)[0]
    assert doc.ready
    doc.code = "SA2026101200078"
    header, details = build_payload(doc)
    rate = Decimal(expected)
    assert details[0]["taxRate"] == rate
    gross, net, tax = money_parts("67.995", rate)
    assert details[0]["origTaxAmount"] == gross
    assert details[0]["origDiscountAmount"] == net
    assert details[0]["origTax"] == tax
    assert header["origAmount"] == net
    assert net + tax == gross
    if rate == Decimal("0.13"):
        assert net == Decimal("60.17") and tax == Decimal("7.825")


@pytest.mark.parametrize("quantity,gross,source_price,net_price", [
    ("10", "67.995", "999", "6.01725664"),
    ("-10", "-67.995", "999", "6.01725664"),
    ("10", "-67.995", "999", "-6.01725664"),
    ("1", "0", "1", "0.00000000"),
    ("0", "0", "6.2", "5.48672566"),
])
def test_amount_based_net_unit_price_and_confirmed_business_flags(quantity, gross, source_price, net_price):
    row = source_row(**{"数量": quantity, "含税金额": gross, "单价": source_price, "业务类型": "16"})
    doc = validate_rows([row], cache(), threading.Event(), lambda *args: None)[0]
    assert doc.ready
    doc.code = "SA2026101200078"
    header, details = build_payload(doc)
    detail = details[0]
    assert detail["saleOutQuantity"] == Decimal(0) and detail["quantity"] == Decimal(quantity)
    assert detail["origDiscountPrice"] == detail["discountPrice"] == Decimal(net_price)
    assert detail["origDiscountPrice"].as_tuple().exponent == -8
    assert detail["origTaxPrice"] == Decimal(source_price)
    assert detail["taxFlag"] == detail["isModifiedPrice"] == 1
    assert detail["RefundDifferenceFlag"] == 51185
    assert header["agency"] == "范立志" and header["agencyid"] == 51
    assert header["isCarriedForwardOut"] == 0
    assert header["invoiceType"] == 281 and header["reciveType"] == 7
    assert header["SaleInvoiceMedia"] == 51014
    assert header["DistributionMode"] == 57419
    assert header["saleInvoiceID"] == 281
    assert header["LogisticsStatus"] == 56134
    for name in ("isCounteractDelivery", "isPriceTrace", "isBeforeSystemInuse",
                 "iscarriedforwardin", "ismodifiedcode", "PrintCount",
                 "IsSeparateByWareHouse", "IsAutoWarehouse", "SendInfoSource",
                 "IsReadOnly", "IsDigElectronicInvoice", "IsOnlinePay"):
        assert header[name] == 0
    for name in ("isPresent", "isManualCost", "IsMemberIntegral", "IsPromotionPresent",
                 "DiscountPriceChange", "TaxPriceChange"):
        assert detail[name] == 0
    assert header["isCancel"] == 428
    assert header["isNoArapBookkeeping"] == 0
    assert "isNoArapBookkeeping" not in detail
    assert all(field not in header for field in ("auditor", "auditorid", "auditeddate", "AuditedTime"))
    assert header["voucherState"] == 181
    for field in ("origprereceiveamount", "PreReceiveAmount", "origSettleAmount", "settleAmount",
                  "OrigInvoiceTaxAmount", "Allowances", "origAllowances", "thisBalanceIntegral"):
        assert type(header[field]) is Decimal and header[field] == Decimal(0)
        assert field not in detail
    assert header["isAutoGenerateInvoice"] == 0
    for payload in (header, detail):
        assert payload["DataSource"] == 2497 and payload["isSaleOut"] == 305
    assert header["origTaxAmount"] == detail["origDiscountAmount"] + detail["origTax"]


@pytest.mark.parametrize("gross,expected_net", [("1.01", "0.51"), ("-1.01", "-0.51")])
def test_net_rounding_half_ties_away_from_zero(gross, expected_net):
    g, n, t = money_parts(gross, Decimal("1"))
    assert n == Decimal(expected_net)
    assert g == n + t


def test_detail_public_and_private_source_fields_match_each_original_line():
    rows = [source_row(" SOURCE-001 ", number=2, **{"销售标识": "LINE-001", "操作员": "源操作员"}),
            source_row(" SOURCE-001 ", number=3, **{"销售标识": "LINE-002", "操作员": "源操作员"})]
    doc = validate_rows(rows, cache(), threading.Event(), lambda *args: None)[0]
    assert doc.ready
    doc.code = "SA2026101200078"
    header, details = build_payload(doc)
    for row, detail in zip(rows, details):
        assert detail["pubuserdefnvc1"] == detail["priuserdefnvc2"] == "SOURCE-001"
        assert detail["pubuserdefnvc2"] == detail["priuserdefnvc3"] == row.values["销售标识"]
    assert header["pubuserdefnvc1"] == ""


@pytest.mark.parametrize("operator", ["  窦秀秀C  ", "", None])
def test_header_ignores_source_operator_without_changing_fixed_users(operator):
    row = source_row(**{"操作员": operator})
    doc = validate_rows([row], cache(), threading.Event(), lambda *args: None)[0]
    assert doc.ready
    header, details = build_payload(doc)
    assert header["pubuserdefnvc1"] == ""
    assert header["maker"] == header["agency"] == "范立志" and "auditor" not in header
    assert header["makerid"] == header["agencyid"] == 51 and "auditorid" not in header
    assert details[0]["pubuserdefnvc1"] == row.voucher


def test_different_source_operators_do_not_block_documents():
    rows = [source_row("A", number=2, **{"操作员": "甲"}),
            source_row("A", number=3, **{"操作员": "乙"}),
            source_row("B", number=4, **{"操作员": "丙"})]
    docs = validate_rows(rows, cache(), threading.Event(), lambda *args: None)
    assert all(doc.ready for doc in docs)
    assert all(not row.reasons for row in rows)
    assert all(build_payload(doc)[0]["pubuserdefnvc1"] == "" for doc in docs)


def test_zero_return_base_unit_project_null_and_quantity_text():
    row = source_row(**{"业务类型": "16", "数量": "0", "含税金额": "-0.019", "项目编码": "999999"})
    doc = validate_rows([row], cache(), threading.Event(), lambda *args: None)[0]
    assert doc.ready
    doc.code = "SA2026101200078"
    header, details = build_payload(doc)
    assert header["idbusinesstype"] == 66
    assert header["idproject"] is None and details[0]["idproject"] is None
    assert details[0]["quantity"] == 0 and details[0]["compositionQuantity"] == "0盒"
    assert details[0]["origTaxAmount"] == Decimal("-0.019")
    assert "auditeddate" not in header
    assert header["isAutoGenerateSaleOut"] == header["IsAutoSaleOut"] == 1
    assert header["saleInvoiceID"] == 281 and "docid" not in header
    assert decimal_text(Decimal("10.0000")) == "10"


def test_missing_customer_blocks_whole_document_but_keeps_all_reasons():
    rows = [source_row(number=2), source_row(number=3, **{"客户编码": "missing", "商品编码": "missing"})]
    docs = validate_rows(rows, cache(), threading.Event(), lambda *args: None)
    assert not docs[0].ready
    assert all(row.status == "校验失败" for row in rows)  # customer header conflict also affects first row
    assert any("商品编码" in reason for reason in rows[1].reasons)
    assert any("客户编码" in reason for reason in rows[1].reasons)


def test_invalid_detail_blocks_healthy_sibling():
    rows = [source_row(number=2), source_row(number=3, **{"商品编码": "missing"})]
    validate_rows(rows, cache(), threading.Event(), lambda *args: None)
    assert rows[0].status == "整单受阻" and rows[1].status == "校验失败"


def test_duplicate_source_skips_even_with_invalid_new_payload(tmp_path):
    repo = FakeRepository()
    repo.committed[key("EXISTING")] = dict(id=42, code="SA2026101200001")
    row = source_row("EXISTING", **{"客户编码": "unknown"})
    result, _ = run_engine(tmp_path, [row], repo)
    assert row.status == "重复跳过" and row.header_id == 42 and row.reasons
    assert not repo.calls and Path(result["report"]).exists()


def test_batch_data_error_isolates_whole_documents(tmp_path):
    repo = FakeRepository()
    repo.fail_source = "BAD"
    rows = [source_row("GOOD1"), source_row("BAD", number=3), source_row("BAD", number=4), source_row("GOOD2")]
    result, _ = run_engine(tmp_path, rows, repo)
    assert [row.status for row in rows] == ["导入成功", "写入失败", "写入失败", "导入成功"]
    assert set(repo.committed) == {"good1", "good2"}
    assert len(repo.calls) > 1 and result["report"]


def test_stop_retains_current_commit_and_exports_remainder(tmp_path):
    stop, repo = threading.Event(), FakeRepository()
    repo.stop_after = stop
    rows = [source_row("FIRST"), source_row("SECOND")]
    result, _ = run_engine(tmp_path, rows, repo, stop, limit=1)
    assert [row.status for row in rows] == ["导入成功", "未处理"]
    assert len(repo.calls) == 1 and Path(result["report"]).exists()


@pytest.mark.parametrize("outcome,status", [("committed", "导入成功"), ("absent", "未处理"), ("corrupt", "待确认")])
def test_lost_commit_response_never_blindly_reinserts(tmp_path, outcome, status):
    repo = FakeRepository()
    repo.unknown = outcome
    row = source_row()
    result, _ = run_engine(tmp_path, [row], repo)
    assert row.status == status and len(repo.calls) == 1
    assert Path(result["report"]).exists()


def test_config_missing_archive_does_not_write(tmp_path):
    repo = FakeRepository()
    repo.preflight_error = "EAP_User.id=51不存在"
    row = source_row()
    result, _ = run_engine(tmp_path, [row], repo)
    assert row.status == "未处理" and not repo.calls
    assert "EAP_User" in result["message"]


def test_rerun_uses_source_not_new_sa_number(tmp_path):
    repo = FakeRepository()
    row = source_row()
    run_engine(tmp_path, [row], repo)
    row.values["含税金额"] = "999"
    run_engine(tmp_path, [row], repo)
    assert row.status == "重复跳过" and len(repo.calls) == 1


def test_report_literal_cells_and_nonblocking_notes(tmp_path):
    rows = [source_row('=HYPERLINK("x")'), source_row("GOOD")]
    rows[0].status = "校验失败"
    rows[0].reasons = ["=1+2"]
    rows[1].status, rows[1].notes = "导入成功", ["项目未匹配留空"]
    report = export_report(rows, list(rows[0].values), tmp_path / "source.xls", "db", "task")
    workbook = load_workbook(report)
    assert workbook["失败及未处理"].cell(2, 3).data_type == "s"
    assert workbook["失败及未处理"].cell(2, 3).value.startswith("=")
    assert workbook["成功但有说明"].max_row == 2
    assert workbook["重复跳过"].max_row == 1
    workbook.close()
    second = export_report(rows, list(rows[0].values), tmp_path / "source.xls", "db", "task")
    assert second != report and report.exists()


def test_journal_roundtrip_without_credentials(tmp_path):
    row = source_row()
    row.status = "待确认"
    row.parsed["gross"] = Decimal("67.99")
    journal = TaskJournal(tmp_path, "task", "source.xls", "db")
    journal.save([row], "写入中")
    payload, rows = TaskJournal.load(journal.path)
    assert rows[0].status == "待确认" and rows[0].parsed["gross"] == "67.99"
    assert "do-not-log-me" not in journal.path.read_text(encoding="utf-8")


def test_config_no_interpolation_and_brace_quoting(tmp_path):
    path = tmp_path / "config.ini"
    path.write_text("[sqlserver]\nip=server\nport=1433\ndatabase=business\nuser=test\npassword=x%y};UID=other\n", encoding="utf-8")
    cfg = DatabaseConfig.load(path)
    assert cfg.password == "x%y};UID=other"
    assert "PWD={x%y}};UID=other}" in cfg.connection_string()
    assert cfg.password not in cfg.display() and cfg.password not in repr(cfg)
    assert cfg.password not in cfg.safe_error(Exception(cfg.connection_string()))


def test_day_overflow_blocks_all_documents_of_that_day(tmp_path):
    repo = FakeRepository()
    repo.max_sequence = lambda day: 99998
    rows = [source_row("A"), source_row("B")]
    run_engine(tmp_path, rows, repo)
    assert all(row.status == "校验失败" for row in rows) and not repo.calls
